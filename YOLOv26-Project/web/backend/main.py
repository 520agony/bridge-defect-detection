"""
桥梁病害检测 Web 后端 (FastAPI + SQLite)
=========================================
功能: 用户注册/登录、单张与批量检测、检测历史、统计分析、数据集信息

启动:
    ./.venv/Scripts/python.exe -m uvicorn web.backend.main:app --host 0.0.0.0 --port 8000
"""

import os
import sys
import time
import uuid
from datetime import datetime
from pathlib import Path

import cv2
import numpy as np
from fastapi import Depends, FastAPI, File, Header, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import Response
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel

from . import auth, db

# ============================================================
# 配置
# ============================================================
PROJECT_ROOT = Path(__file__).resolve().parent.parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))
MODEL_PATH = Path(os.environ.get("MODEL_PATH", PROJECT_ROOT / "weights" / "bridge_defect_best.pt"))
# 融合演示专用权重（仿真场景按 distill 调优；缺失则回退主权重）
FUSION_MODEL_PATH = Path(os.environ.get("FUSION_MODEL_PATH", r"D:\开始做项目\第三轮\best.pt"))
_fusion_model = None


def get_fusion_model():
    """融合演示用模型（懒加载 + 预热）"""
    global _fusion_model
    if _fusion_model is None:
        from ultralytics import YOLO
        p = FUSION_MODEL_PATH if FUSION_MODEL_PATH.exists() else MODEL_PATH
        _fusion_model = YOLO(str(p))
        try:
            _fusion_model.predict(np.zeros((64, 64, 3), dtype=np.uint8), imgsz=IMGSZ, verbose=False)
        except Exception:
            pass
    return _fusion_model
CONF_THRESHOLD = float(os.environ.get("CONF_THRESHOLD", "0.25"))
IMGSZ = int(os.environ.get("IMGSZ", "640"))
UPLOAD_DIR = Path(__file__).resolve().parent / "data" / "images"
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
VIDEO_DIR = Path(__file__).resolve().parent / "data" / "videos"
VIDEO_DIR.mkdir(parents=True, exist_ok=True)
MAX_VIDEO_FRAMES = int(os.environ.get("MAX_VIDEO_FRAMES", "1800"))  # 视频检测帧上限（约 60s@30fps）
VIDEO_BATCH = int(os.environ.get("VIDEO_BATCH", "8"))  # 视频批推理帧数（提速用）

_model = None


def get_model():
    """懒加载 YOLO 模型（含一次预热推理，消除首帧 CUDA 初始化开销）"""
    global _model
    if _model is None:
        from ultralytics import YOLO
        if not MODEL_PATH.exists():
            raise FileNotFoundError(f"模型权重不存在: {MODEL_PATH}")
        _model = YOLO(str(MODEL_PATH))
        try:
            # 预热两种推理形状：单帧（摄像头）与批 8（视频），消除首调 CUDA 初始化开销
            _model.predict(np.zeros((64, 64, 3), dtype=np.uint8), imgsz=IMGSZ, verbose=False)
            _model.predict([np.zeros((64, 64, 3), dtype=np.uint8)] * VIDEO_BATCH, imgsz=IMGSZ, verbose=False)
        except Exception:
            pass
    return _model


def _imwrite(path, img, params=None):
    """cv2.imwrite 的中文路径兼容版（imencode + numpy tofile）"""
    ok, buf = cv2.imencode(Path(path).suffix or ".jpg", img, params or [])
    if ok:
        buf.tofile(str(path))
    return ok


# ============================================================
# 数据集信息（5 类，静态统计）
# ============================================================
DATASET_CLASSES = [
    {"id": 0, "name": "crack",          "zh": "裂缝", "boxes": 21004, "color": "#e6194b"},
    {"id": 1, "name": "spalling",       "zh": "剥落", "boxes": 12653, "color": "#f58231"},
    {"id": 2, "name": "hole",           "zh": "孔洞", "boxes": 6501,  "color": "#3cb44b"},
    {"id": 3, "name": "rebar_exposed",  "zh": "露筋", "boxes": 16594, "color": "#42d4f4"},
    {"id": 4, "name": "efflorescence", "zh": "泛碱", "boxes": 18102, "color": "#a78bfa"},
]
DATASET_SPLITS = {"train": 30216, "val": 4822, "test": 2406}


app = FastAPI(title="桥梁病害检测平台", version="2.0.0")
app.add_middleware(CORSMiddleware, allow_origins=["*"], allow_credentials=True,
                   allow_methods=["*"], allow_headers=["*"])


@app.on_event("startup")
def _startup_warm():
    """服务启动后在后台线程预热模型（首次推理请求不再等待加载）"""
    import threading

    def _warm():
        try:
            get_model()
        except Exception:
            pass

    threading.Thread(target=_warm, daemon=True).start()


@app.on_event("startup")
def _startup_cleanup():
    """清理 7 天前的输出视频/热力图与上传临时文件（检测历史图片保留）"""
    import time as _time

    now = _time.time()
    try:
        for d, pats in ((VIDEO_DIR, ("out_*", "in_*")), (UPLOAD_DIR, ("*_heatmap.jpg",))):
            if not d.exists():
                continue
            for pat in pats:
                for p in d.glob(pat):
                    try:
                        if now - p.stat().st_mtime > 7 * 86400:
                            p.unlink()
                    except Exception:
                        pass
    except Exception:
        pass

    # 打印局域网二维码（手机扫码直达）
    try:
        import socket as _socket

        ip = _socket.gethostbyname(_socket.gethostname())
        if ip.startswith("127."):
            import subprocess as _sp

            out = _sp.run(["ipconfig"], capture_output=True, text=True, timeout=10).stdout
            import re as _re

            m = _re.search(r"IPv4[^:]*:\s*(10\.\d+\.\d+\.\d+|192\.168\.\d+\.\d+|172\.(1[6-9]|2\d|3[01])\.\d+\.\d+)",
                          out)
            if m:
                ip = m.group(1)
        url = f"http://{ip}:8000/"
        print(f"\n[前端地址] 本机 http://127.0.0.1:8000/ ｜ 局域网 {url}")
        try:
            import qrcode as _qr

            _qr.QRCode(border=1).add_data(url).print_ascii(invert=True)
        except Exception:
            pass
    except Exception:
        pass


# ============================================================
# 认证
# ============================================================
class RegisterBody(BaseModel):
    username: str
    password: str
    email: str


class LoginBody(BaseModel):
    username: str
    password: str


@app.post("/auth/register")
def register(body: RegisterBody):
    err = auth.validate_register(body.username, body.password, body.email)
    if err:
        raise HTTPException(status_code=400, detail=err)

    conn = db.get_db()
    exists = conn.execute("SELECT id FROM users WHERE username = ?", (body.username,)).fetchone()
    if exists:
        conn.close()
        raise HTTPException(status_code=409, detail="用户名已被占用")

    salt = uuid.uuid4().hex
    phash = auth.hash_password(body.password, salt)
    cur = conn.execute(
        "INSERT INTO users (username, email, password_hash, salt, created_at) VALUES (?,?,?,?,?)",
        (body.username, body.email, phash, salt, datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
    )
    user_id = cur.lastrowid
    token = auth.create_token()
    conn.execute("INSERT INTO sessions (token, user_id, created_at) VALUES (?,?,?)",
                 (token, user_id, datetime.now().strftime("%Y-%m-%d %H:%M:%S")))
    conn.commit()
    conn.close()
    return {"token": token, "username": body.username, "email": body.email}


@app.post("/auth/login")
def login(body: LoginBody):
    conn = db.get_db()
    row = conn.execute("SELECT * FROM users WHERE username = ?", (body.username,)).fetchone()
    if row is None or not auth.verify_password(body.password, row["salt"], row["password_hash"]):
        conn.close()
        raise HTTPException(status_code=401, detail="用户名或密码错误")

    token = auth.create_token()
    conn.execute("INSERT INTO sessions (token, user_id, created_at) VALUES (?,?,?)",
                 (token, row["id"], datetime.now().strftime("%Y-%m-%d %H:%M:%S")))
    conn.commit()
    conn.close()
    return {"token": token, "username": row["username"], "email": row["email"]}


@app.post("/auth/logout")
def logout(authorization: str = Header(default=""), user: dict = Depends(auth.get_current_user)):
    """注销：删除服务端 session，使 token 真正失效"""
    token = authorization.removeprefix("Bearer ").strip() if authorization.startswith("Bearer ") else ""
    if token:
        conn = db.get_db()
        conn.execute("DELETE FROM sessions WHERE token = ?", (token,))
        conn.commit()
        conn.close()
    return {"status": "ok"}


@app.get("/auth/me")
def me(user: dict = Depends(auth.get_current_user)):
    return user


# ============================================================
# 检测
# ============================================================
def _run_detect(img: np.ndarray, conf: float = CONF_THRESHOLD, _model_override=None):
    """对一张图推理，返回 (detections 列表, inference_ms)。
    每项含 bbox + mask 多边形点列 + 量化测量 + 等级评估。
    """
    from src.analysis.measure import measure_defect
    from src.analysis.grade import grade_defect

    model = _model_override or get_model()
    h, w = img.shape[:2]
    results = model.predict(img, conf=conf, imgsz=IMGSZ, verbose=False)
    r = results[0]
    detections = []
    if r.masks is None:
        return detections, round(float(r.speed['inference']), 2)
    # 归一化 mask 坐标转回原图绝对坐标
    h_ratio, w_ratio = h / IMGSZ, w / IMGSZ
    for i, box in enumerate(r.boxes):
        x1, y1, x2, y2 = box.xyxy[0].tolist()
        cls_id = int(box.cls[0])
        cls_name = r.names[cls_id]
        conf = float(box.conf[0])
        # 取 mask 多边形（xy 归一化 → 原图绝对像素）
        poly = r.masks.xyn[i].reshape(-1, 2).tolist()
        poly_abs = [[round(p[0] * w, 1), round(p[1] * h, 1)] for p in poly]
        bbox = [x1, y1, x2, y2]
        try:
            meas = measure_defect(cls_name, img, bbox, mask=poly_abs)
            grade = grade_defect(cls_name, meas)
        except Exception:
            meas, grade = {}, {"grade": 0, "name": "未知", "action": "人工判定", "needs_calibration": True}
        detections.append({
            "class_id": cls_id,
            "class": cls_name,
            "confidence": round(conf, 4),
            "bbox": [round(x1, 1), round(y1, 1), round(x2, 1), round(y2, 1)],
            "mask": poly_abs,  # 原图像素坐标
            "mask_points": len(poly),  # 多边形顶点数
            "measurements": meas,
            "grade": grade,
        })
    detections.sort(key=lambda d: d["confidence"], reverse=True)
    _append_composite_hints(img, detections)
    return detections, round(float(r.speed["inference"]), 2)


def _bbox_iou(a, b) -> float:
    x1, y1 = max(a[0], b[0]), max(a[1], b[1])
    x2, y2 = min(a[2], b[2]), min(a[3], b[3])
    inter = max(0.0, x2 - x1) * max(0.0, y2 - y1)
    if inter <= 0:
        return 0.0
    aa = max(0.0, a[2] - a[0]) * max(0.0, a[3] - a[1])
    bb = max(0.0, b[2] - b[0]) * max(0.0, b[3] - b[1])
    return inter / max(1e-9, aa + bb - inter)


def _append_composite_hints(img: np.ndarray, detections: list):
    """复合病害提示（规则层）：大剥落内无模型露筋检出时，跑锈色线状检测。
    命中则追加 rule_hint 条目（虚线框 + 徽章，不计入成绩）。就地修改 detections。"""
    try:
        from src.analysis.composite import SPALL_FRAC_MIN, suggest_rebar_in_spall
    except Exception:
        return
    h, w = img.shape[:2]
    rebar_boxes = [d["bbox"] for d in detections
                   if d.get("class") == "rebar_exposed" and not d.get("rule_hint")]
    for d in list(detections):
        if d.get("class") != "spalling" or d.get("rule_hint") or not d.get("mask"):
            continue
        area_px = (d.get("measurements") or {}).get("area_px") or 0
        if area_px / max(1, h * w) < SPALL_FRAC_MIN:
            continue
        if any(_bbox_iou(d["bbox"], rb) > 0.1 for rb in rebar_boxes):
            continue
        try:
            hit, info = suggest_rebar_in_spall(img, d["mask"], d["bbox"])
        except Exception:
            continue
        if not hit:
            continue
        rb = info.get("rust_bbox") or d["bbox"]
        score = 0.55 + 0.4 * min(1.0, (info.get("rust_ratio", 0) / 0.3
                                       + info.get("line_density", 0) / 30) / 2)
        detections.append({
            "class_id": 3,
            "class": "rebar_exposed",
            "confidence": round(min(0.95, score), 4),
            "bbox": [round(float(v), 1) for v in rb],
            "mask": [],
            "mask_points": 0,
            "measurements": {},
            "grade": {},
            "rule_hint": True,
            "rule": "spall_rust_lines",
            "rule_info": info,
        })
    detections.sort(key=lambda d: d["confidence"], reverse=True)


def _save_detection(user_id: int, filename: str, img: np.ndarray, detections: list) -> dict:
    """保存检测记录（图片落盘 + 结果入库；count 只计模型检出，规则提示不计）"""
    stored_name = f"{uuid.uuid4().hex}_{filename}"
    stored_path = UPLOAD_DIR / stored_name
    _imwrite(stored_path, img, [cv2.IMWRITE_JPEG_QUALITY, 90])

    result_json = __import__("json").dumps(detections, ensure_ascii=False)
    model_count = sum(1 for d in detections if not d.get("rule_hint"))
    conn = db.get_db()
    cur = conn.execute(
        "INSERT INTO detections (user_id, filename, image_path, result_json, count, created_at) "
        "VALUES (?,?,?,?,?,?)",
        (user_id, filename, stored_name, result_json, model_count,
         datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
    )
    det_id = cur.lastrowid
    conn.commit()
    conn.close()
    return {"id": det_id, "filename": filename, "image": stored_name,
            "count": model_count, "detections": detections}


@app.post("/detect")
async def detect(file: UploadFile = File(...), user: dict = Depends(auth.get_current_user)):
    img_bytes = await file.read()
    img = cv2.imdecode(np.frombuffer(img_bytes, np.uint8), cv2.IMREAD_COLOR)
    if img is None:
        raise HTTPException(status_code=400, detail="无法解析图片")

    try:
        detections, ms = _run_detect(img)
    except FileNotFoundError as e:
        raise HTTPException(status_code=503, detail=str(e))

    rec = _save_detection(user["id"], file.filename or "image.jpg", img, detections)
    h, w = img.shape[:2]
    return {"id": rec["id"], "detections": detections, "count": rec["count"],
            "image_size": [w, h], "inference_ms": ms,
            "classes": {c["id"]: c["name"] for c in DATASET_CLASSES}}


# ============================================================
# 摄像头实时（轻量帧检测：不入库、不落盘）
# ============================================================
@app.post("/detect/frame")
async def detect_frame(file: UploadFile = File(...), user: dict = Depends(auth.get_current_user)):
    """单帧推理（供前端摄像头实时检测循环调用）"""
    img_bytes = await file.read()
    img = cv2.imdecode(np.frombuffer(img_bytes, np.uint8), cv2.IMREAD_COLOR)
    if img is None:
        raise HTTPException(status_code=400, detail="无法解析图片")
    try:
        detections, ms = _run_detect(img)
    except FileNotFoundError as e:
        raise HTTPException(status_code=503, detail=str(e))
    h, w = img.shape[:2]
    return {"detections": detections, "count": len(detections),
            "image_size": [w, h], "inference_ms": ms,
            "classes": {c["id"]: c["name"] for c in DATASET_CLASSES}}


# ============================================================
# 视频检测（上传 → 逐帧推理 → 输出标注视频 + 统计，不入检测历史）
# ============================================================
@app.post("/detect/video")
async def detect_video(file: UploadFile = File(...), user: dict = Depends(auth.get_current_user)):
    """整段视频逐帧检测并生成标注视频（上限 MAX_VIDEO_FRAMES 帧）"""
    raw = await file.read()
    suffix = Path(file.filename or "v.mp4").suffix or ".mp4"
    return _detect_video_bytes(raw, suffix)


def _detect_video_bytes(raw: bytes, suffix: str) -> dict:
    """视频字节 → 逐帧推理 → 标注视频 + 统计（供上传/素材库复用）"""
    if len(raw) > 300 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="视频过大（>300MB），请先压缩或截取")
    tmp_in = VIDEO_DIR / f"in_{uuid.uuid4().hex}{suffix}"
    tmp_in.write_bytes(raw)
    cap = cv2.VideoCapture(str(tmp_in))
    if not cap.isOpened():
        tmp_in.unlink(missing_ok=True)
        raise HTTPException(status_code=400, detail="无法解析视频（支持 mp4/avi/mov/webm）")

    fps = float(cap.get(cv2.CAP_PROP_FPS) or 25.0)
    total = int(cap.get(cv2.CAP_PROP_FRAME_COUNT) or 0)
    vw = int(cap.get(cv2.CAP_PROP_FRAME_WIDTH))
    vh = int(cap.get(cv2.CAP_PROP_FRAME_HEIGHT))
    if vw <= 0 or vh <= 0:
        cap.release(); tmp_in.unlink(missing_ok=True)
        raise HTTPException(status_code=400, detail="视频尺寸异常")
    limit = min(total if total > 0 else MAX_VIDEO_FRAMES, MAX_VIDEO_FRAMES)

    out_name = f"out_{uuid.uuid4().hex}.mp4"
    raw_name = out_name.replace(".mp4", "_raw.avi")
    writer = cv2.VideoWriter(str(VIDEO_DIR / raw_name), cv2.VideoWriter_fourcc(*"XVID"), fps, (vw, vh))
    if not writer.isOpened():
        raw_name = out_name.replace(".mp4", "_raw.mp4")
        writer = cv2.VideoWriter(str(VIDEO_DIR / raw_name), cv2.VideoWriter_fourcc(*"mp4v"), fps, (vw, vh))
    if not writer.isOpened():
        cap.release(); tmp_in.unlink(missing_ok=True)
        raise HTTPException(status_code=500, detail="无法创建输出视频（编码器不可用）")

    try:
        model = get_model()
    except FileNotFoundError as e:
        cap.release(); writer.release(); tmp_in.unlink(missing_ok=True)
        raise HTTPException(status_code=503, detail=str(e))

    class_counts: dict = {}
    processed = 0
    t0 = time.time()
    try:
        while processed < limit:
            batch = []
            while len(batch) < VIDEO_BATCH and processed + len(batch) < limit:
                ok, frame = cap.read()
                if not ok:
                    break
                batch.append(frame)
            if not batch:
                break
            results = model.predict(batch, conf=CONF_THRESHOLD, imgsz=IMGSZ, verbose=False)
            for r in results:
                for b in r.boxes:
                    cn = r.names[int(b.cls[0])]
                    class_counts[cn] = class_counts.get(cn, 0) + 1
                writer.write(r.plot())
            processed += len(batch)
    finally:
        cap.release()
        writer.release()
        tmp_in.unlink(missing_ok=True)

    if processed == 0:
        (VIDEO_DIR / raw_name).unlink(missing_ok=True)
        raise HTTPException(status_code=400, detail="视频没有可读取的帧")

    # 转 H.264（浏览器可在线播）+ faststart（边下边播）；失败则回退原文件（仅可下载）
    h264_ok = False
    try:
        import shutil as _shutil
        import subprocess as _sp
        ff = _shutil.which("ffmpeg")
        if ff:
            _sp.run([ff, "-y", "-i", str(VIDEO_DIR / raw_name), "-c:v", "libx264",
                     "-preset", "veryfast", "-crf", "23", "-pix_fmt", "yuv420p",
                     "-movflags", "+faststart", "-an", str(VIDEO_DIR / out_name)],
                    check=True, capture_output=True, timeout=600)
            (VIDEO_DIR / raw_name).unlink(missing_ok=True)
            h264_ok = True
    except Exception:
        h264_ok = False
    if not h264_ok:
        (VIDEO_DIR / out_name).unlink(missing_ok=True)
        out_name = raw_name

    elapsed = time.time() - t0
    return {
        "video": out_name,
        "preview": h264_ok,
        "frames": processed,
        "fps": round(fps, 1),
        "duration_s": round(processed / fps, 1) if fps else None,
        "truncated": bool(total > 0 and total > limit),
        "max_frames": MAX_VIDEO_FRAMES,
        "elapsed_s": round(elapsed, 1),
        "frame_ms_avg": round(elapsed * 1000 / processed, 1),
        "class_counts": class_counts,
        "classes": {c["id"]: c["name"] for c in DATASET_CLASSES},
    }


# ============================================================
# Grad-CAM 热力图
# ============================================================
def _generate_heatmap(model, img_bgr):
    """生成 Grad-CAM 热力图叠加图，返回 (BGR 图, score)"""
    import torch
    from src.eval.gradcam import grad_cam, overlay_heatmap

    detect = model.model
    device = next(detect.parameters()).device  # 模型所在设备
    img_rgb = cv2.cvtColor(img_bgr, cv2.COLOR_BGR2RGB)
    img_r = cv2.resize(img_rgb, (640, 640))
    x = torch.from_numpy(img_r).permute(2, 0, 1).float() / 255.0
    x = x.unsqueeze(0).to(device)
    cam, score = grad_cam(model, x, -2)
    return overlay_heatmap(img_r, cam), float(score)


@app.post("/heatmap")
async def heatmap(file: UploadFile = File(...), user: dict = Depends(auth.get_current_user)):
    img_bytes = await file.read()
    img = cv2.imdecode(np.frombuffer(img_bytes, np.uint8), cv2.IMREAD_COLOR)
    if img is None:
        raise HTTPException(status_code=400, detail="无法解析图片")

    try:
        model = get_model()
        heat_img, score = _generate_heatmap(model, img)
    except FileNotFoundError as e:
        raise HTTPException(status_code=503, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"热力图生成失败: {e}")

    stored_name = f"{uuid.uuid4().hex}_heatmap.jpg"
    _imwrite(UPLOAD_DIR / stored_name, heat_img, [cv2.IMWRITE_JPEG_QUALITY, 92])
    return {"image": stored_name, "score": round(score, 4)}


# ============================================================
# 点云后融合演示（B2：框/mask→投影筛点→米制 vs 真值）
# ============================================================
FUSION_DIR = Path(__file__).resolve().parent / "data" / "fusion_samples"


def _fusion_sample_list():
    """扫描融合场景目录（rgb.png + cloud.ply + meta.json + truth.json 齐全才收录）"""
    out = []
    if FUSION_DIR.exists():
        for d in sorted(FUSION_DIR.iterdir()):
            if not d.is_dir():
                continue
            if not all((d / f).exists() for f in ("rgb.png", "cloud.ply", "meta.json", "truth.json")):
                continue
            try:
                truth = __import__("json").loads((d / "truth.json").read_text(encoding="utf-8"))
            except Exception:
                truth = {}
            out.append({"id": d.name, "truth_keys": sorted(truth.keys())})
    return out


FUSION_SAMPLES = tuple(s["id"] for s in _fusion_sample_list()) or ("sample03", "sample02")


@app.get("/fusion/samples")
def fusion_samples(user: dict = Depends(auth.get_current_user)):
    return {"samples": _fusion_sample_list(), "count": len(FUSION_SAMPLES)}


# ============================================================
# 演示素材库（图片 meta.json 描述类别 + 视频目录扫描）
# ============================================================
LIB_DIR = Path(__file__).resolve().parent / "data" / "lib"
LIB_VIDEO_DIR = Path(__file__).resolve().parent / "data" / "lib_videos"
LIB_VIDEO_DIR.mkdir(parents=True, exist_ok=True)
VIDEO_SUFFIXES = (".mp4", ".avi", ".mov", ".webm", ".mkv")


@app.get("/lib")
def lib_list(user: dict = Depends(auth.get_current_user)):
    try:
        meta = __import__("json").loads((LIB_DIR / "meta.json").read_text(encoding="utf-8"))
    except Exception:
        meta = []
    videos = []
    vmeta = {}
    try:
        for m in __import__("json").loads((LIB_VIDEO_DIR / "meta.json").read_text(encoding="utf-8")):
            vmeta[m["file"]] = m
    except Exception:
        pass
    for p in sorted(LIB_VIDEO_DIR.iterdir()):
        if p.suffix.lower() in VIDEO_SUFFIXES:
            m = vmeta.get(p.name, {})
            videos.append({"file": p.name, "mb": round(p.stat().st_size / 1048576, 1),
                           "classes": m.get("classes", []), "zh": m.get("zh", [])})
    return {"images": meta, "count": len(meta), "videos": videos, "video_count": len(videos)}


class LibDetectBody(BaseModel):
    file: str = ""


@app.post("/lib/detect")
def lib_detect(body: LibDetectBody, user: dict = Depends(auth.get_current_user)):
    """素材库图片直接检测（入库，可导出报告）"""
    name = Path(body.file or "").name
    if not name:
        raise HTTPException(status_code=400, detail="缺少文件名")
    p = LIB_DIR / name
    if not p.exists():
        raise HTTPException(status_code=404, detail="素材不存在")
    img = cv2.imdecode(np.fromfile(str(p), dtype=np.uint8), cv2.IMREAD_COLOR)
    if img is None:
        raise HTTPException(status_code=400, detail="素材图片损坏")
    try:
        detections, ms = _run_detect(img)
    except FileNotFoundError as e:
        raise HTTPException(status_code=503, detail=str(e))
    rec = _save_detection(user["id"], name, img, detections)
    h, w = img.shape[:2]
    return {"id": rec["id"], "detections": detections, "count": rec["count"],
            "image": rec["image"], "image_size": [w, h], "inference_ms": ms,
            "classes": {c["id"]: c["name"] for c in DATASET_CLASSES}}


@app.post("/lib/detect_video")
def lib_detect_video(body: LibDetectBody, user: dict = Depends(auth.get_current_user)):
    """素材库视频直接检测（不入库，返回标注视频+统计）"""
    name = Path(body.file or "").name
    p = LIB_VIDEO_DIR / name
    if not name or not p.exists():
        raise HTTPException(status_code=404, detail="素材视频不存在")
    return _detect_video_bytes(p.read_bytes(), p.suffix or ".mp4")


class FusionBody(BaseModel):
    sample: str = "sample03"
    conf: float = 0.25


@app.post("/fusion/demo")
def fusion_demo(body: FusionBody, user: dict = Depends(auth.get_current_user)):
    from src.fusion.late_fusion import (load_camera, read_ply_ascii, points_in_box,
        points_in_mask, statistical_outlier_filter, measure_length_3d,
        measure_area_3d, median_depth)

    if body.sample not in FUSION_SAMPLES:
        raise HTTPException(status_code=400, detail="未知场景")
    sdir = FUSION_DIR / body.sample
    rgb_p = sdir / "rgb.png"
    if not rgb_p.exists():
        raise HTTPException(status_code=404, detail="场景缺文件")
    img = cv2.imdecode(np.frombuffer(rgb_p.read_bytes(), np.uint8), cv2.IMREAD_COLOR)
    if img is None:
        raise HTTPException(status_code=500, detail="场景图片损坏")
    h, w = img.shape[:2]
    K, R, t, W, H = load_camera(str(sdir / "meta.json"))
    P, _ = read_ply_ascii(str(sdir / "cloud.ply"))
    truth = __import__("json").loads((sdir / "truth.json").read_text(encoding="utf-8"))

    detections, ms = _run_detect(img, conf=max(0.05, min(0.8, body.conf)),
                                 _model_override=get_fusion_model())
    fused = []
    for d in detections:
        if d.get("rule_hint"):
            continue  # 规则提示无真值可比，跳过融合
        idx = points_in_mask(P, K, R, t, d["mask"], W, H) if d.get("mask") \
            else points_in_box(P, K, R, t, d["bbox"], W, H)
        Pc = P[idx]
        keep = statistical_outlier_filter(Pc) if len(Pc) > 8 else np.arange(len(Pc))
        Pf = Pc[keep]
        e = {"class": d["class"], "confidence": d["confidence"],
             "n_raw": int(len(idx)), "n_clean": int(len(Pf))}
        if len(Pf) >= 3:
            e["depth_m"] = round(float(median_depth(Pf, K, R, t)), 3)
            if d["class"] == "crack":
                L = float(measure_length_3d(Pf))
                e["fusion_m"] = round(L, 4)
                e["truth_m"] = truth["crack_length_m"]
                e["err_pct"] = round(abs(L - truth["crack_length_m"]) / truth["crack_length_m"] * 100, 2)
            elif d["class"] == "spalling":
                a = float(measure_area_3d(Pf))
                e["fusion_sqm"] = round(a, 6)
                e["truth_sqm"] = truth["spall_area_sqm"]
                e["err_pct"] = round(abs(a - truth["spall_area_sqm"]) / truth["spall_area_sqm"] * 100, 2)
            else:  # hole 等面积类：有真值则对比误差，否则只给面积
                a = float(measure_area_3d(Pf))
                e["fusion_sqm"] = round(a, 6)
                tv = truth.get("hole_area_sqm")
                if tv:
                    e["truth_sqm"] = tv
                    e["err_pct"] = round(abs(a - tv) / tv * 100, 2)
        else:
            e["note"] = "too_few_points"
        fused.append(e)
    return {"sample": body.sample, "image": f"/fusion_img/{body.sample}/rgb.png",
            "image_size": [w, h], "inference_ms": ms,
            "detections": detections, "fused": fused, "truth": truth,
            "demo_weight": "distill（第三轮）" if FUSION_MODEL_PATH.exists() else "主权重（回退）"}


# ============================================================
# 检测历史
# ============================================================
@app.get("/history")
def history(user: dict = Depends(auth.get_current_user)):
    conn = db.get_db()
    rows = conn.execute(
        "SELECT id, filename, image_path, count, created_at FROM detections "
        "WHERE user_id = ? ORDER BY id DESC LIMIT 200", (user["id"],)).fetchall()
    conn.close()
    return {"history": [dict(r) for r in rows]}


@app.get("/history/{det_id}")
def history_detail(det_id: int, user: dict = Depends(auth.get_current_user)):
    conn = db.get_db()
    row = conn.execute("SELECT * FROM detections WHERE id = ? AND user_id = ?",
                       (det_id, user["id"])).fetchone()
    conn.close()
    if row is None:
        raise HTTPException(status_code=404, detail="记录不存在")
    return {"id": row["id"], "filename": row["filename"], "image": row["image_path"],
            "count": row["count"], "created_at": row["created_at"],
            "detections": __import__("json").loads(row["result_json"])}


@app.delete("/history/{det_id}")
def history_delete(det_id: int, user: dict = Depends(auth.get_current_user)):
    conn = db.get_db()
    row = conn.execute("SELECT image_path FROM detections WHERE id = ? AND user_id = ?",
                       (det_id, user["id"])).fetchone()
    if row is None:
        conn.close()
        raise HTTPException(status_code=404, detail="记录不存在")
    conn.execute("DELETE FROM detections WHERE id = ?", (det_id,))
    conn.commit()
    conn.close()
    (UPLOAD_DIR / row["image_path"]).unlink(missing_ok=True)
    return {"status": "ok"}


# ============================================================
# 检测报告导出（Word / PDF）
# ============================================================
REPORT_ZH = {c["name"]: c["zh"] for c in DATASET_CLASSES}
REPORT_BGR = {
    "crack": (75, 25, 230), "spalling": (49, 130, 245), "hole": (75, 180, 60),
    "rebar_exposed": (244, 212, 66), "efflorescence": (250, 139, 167),
}
CJK_FONT = r"C:\Windows\Fonts\msyh.ttc"


def _measure_text(d: dict) -> str:
    m = d.get("measurements") or {}
    if d.get("class") == "crack" and m.get("length_px") is not None:
        s = f"长{m['length_px']}px·宽{m.get('width_px', '-')}px"
        return s + ("（估算）" if m.get("estimated") else "")
    if d.get("class") == "hole" and m.get("diameter_px") is not None:
        return f"直径约{m['diameter_px']}px"
    if m.get("area_px") is not None:
        s = f"面积约{round(m['area_px'])}px"
        if m.get("area_frac") is not None:
            s += f"（占比{m['area_frac'] * 100:.2f}%）"
        return s
    return "-"


def _grade_text(d: dict) -> str:
    g = d.get("grade") or {}
    if not g.get("grade"):
        return "-"
    s = f"{g.get('grade', '-')}级·{g.get('name', '')}"
    b = (g.get("basis") or {})
    if b.get("short"):
        s += f"（{b['short']}）"
    return s


def _annotate_report(img: np.ndarray, detections: list) -> np.ndarray:
    """画 mask 填充 + 描边 + 中文标签条，返回 BGR 图"""
    from PIL import Image as _PI, ImageDraw as _PD, ImageFont as _PF

    out = img.copy()
    h, w = out.shape[:2]
    overlay = out.copy()
    for d in detections:
        color = REPORT_BGR.get(d.get("class"), (136, 136, 136))
        poly = d.get("mask") or []
        if len(poly) > 2:
            pts = np.array(poly, dtype=np.int32).reshape(-1, 1, 2)
            cv2.fillPoly(overlay, [pts], color)
            cv2.polylines(out, [pts], True, color, max(2, w // 400))
        else:
            x1, y1, x2, y2 = [int(v) for v in d["bbox"]]
            cv2.rectangle(out, (x1, y1), (x2, y2), color, max(2, w // 300))
    cv2.addWeighted(overlay, 0.3, out, 0.7, 0, out)
    # 中文标签（PIL）
    pil = _PI.fromarray(cv2.cvtColor(out, cv2.COLOR_BGR2RGB))
    dr = _PD.Draw(pil)
    try:
        font = _PF.truetype(CJK_FONT, max(14, w // 45))
    except Exception:
        font = _PF.load_default()
    for d in detections:
        color = REPORT_BGR.get(d.get("class"), (136, 136, 136))
        rgb = (color[2], color[1], color[0])
        x1, y1, x2, y2 = [int(v) for v in d["bbox"]]
        label = f"{REPORT_ZH.get(d.get('class'), d.get('class'))} {d.get('confidence', 0) * 100:.0f}%"
        tb = dr.textbbox((0, 0), label, font=font)
        tw, th = tb[2] - tb[0] + 12, tb[3] - tb[1] + 10
        ly = y1 - th if y1 - th > 0 else y1
        dr.rectangle([x1, ly, x1 + tw, ly + th], fill=rgb)
        dr.text((x1 + 6, ly + 4), label, font=font, fill=(10, 13, 18))
    return cv2.cvtColor(np.array(pil), cv2.COLOR_RGB2BGR)


def _report_rows(detections: list):
    rows = []
    for i, d in enumerate(detections, 1):
        cls = REPORT_ZH.get(d.get("class"), d.get("class", "-"))
        if d.get("rule_hint"):
            cls += "（规则提示，非模型检出）"
        rows.append([str(i), cls,
                     f"{d.get('confidence', 0) * 100:.1f}%", _measure_text(d), _grade_text(d)])
    return rows


def _export_docx(filename: str, created_at: str, detections: list, ann_jpg: bytes) -> bytes:
    import io as _io
    from docx import Document as _Doc
    from docx.oxml import OxmlElement as _OE
    from docx.oxml.ns import qn as _qn
    from docx.shared import Pt as _Pt, Cm as _Cm

    doc = _Doc()
    doc.styles["Normal"].font.size = _Pt(10.5)

    doc.add_heading("桥梁病害检测报告", level=1)
    doc.add_paragraph(f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}　·　桥梁病害检测平台")
    doc.add_heading("一、基本信息", level=2)
    t = doc.add_table(rows=1, cols=2)
    t.style = "Table Grid"
    for k, v in [("文件名", filename), ("检测时间", created_at),
                 ("检出病害", f"{len(detections)} 处"),
                 ("模型", "YOLOv26-seg（桥梁病害 5 类）")]:
        row = t.add_row().cells
        row[0].text, row[1].text = k, v
    doc.add_heading("二、标注结果图", level=2)
    doc.add_picture(_io.BytesIO(ann_jpg), width=_Cm(14))
    doc.add_heading("三、病害明细", level=2)
    if detections:
        h = ["序号", "类别", "置信度", "量化测量", "病害等级"]
        t2 = doc.add_table(rows=1, cols=len(h))
        t2.style = "Table Grid"
        for cell, txt in zip(t2.rows[0].cells, h):
            cell.text = txt
        for r in _report_rows(detections):
            for cell, txt in zip(t2.add_row().cells, r):
                cell.text = txt
    else:
        doc.add_paragraph("本次检测未发现病害。")
    doc.add_paragraph("注：量化测量为像素级估算（未做米制标定）；等级依据内置评估规则，仅供参考。")
    # 全文 run 显式指定中文字体（ascii/hAnsi/eastAsia），脱离样式继承，保证任何 Word 正常显示中文
    for p in list(doc.paragraphs) + [p for t in doc.tables for r in t.rows for c in r.cells for p in c.paragraphs]:
        for run in p.runs:
            rpr = run._element.get_or_add_rPr()
            rf = rpr.find(_qn("w:rFonts"))
            if rf is None:
                rf = _OE("w:rFonts")
                rpr.append(rf)
            for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                rf.set(_qn(attr), "微软雅黑")
    buf = _io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _export_pdf(filename: str, created_at: str, detections: list, ann_jpg: bytes) -> bytes:
    import io as _io
    from reportlab.lib import colors as _colors
    from reportlab.lib.pagesizes import A4 as _A4
    from reportlab.lib.styles import ParagraphStyle as _PS
    from reportlab.lib.units import cm as _cm
    from reportlab.pdfbase import pdfmetrics as _pm
    from reportlab.pdfbase.ttfonts import TTFont as _TTF
    from reportlab.platypus import Image as _Img, Paragraph as _P, SimpleDocTemplate as _Doc
    from reportlab.platypus import Spacer as _Sp, Table as _T, TableStyle as _TS

    _pm.registerFont(_TTF("MSYH", CJK_FONT))
    st_title = _PS("t", fontName="MSYH", fontSize=18, leading=24, alignment=1, spaceAfter=12)
    st_h = _PS("h", fontName="MSYH", fontSize=13, leading=18, spaceBefore=14, spaceAfter=8)
    st_n = _PS("n", fontName="MSYH", fontSize=10, leading=15)
    st_c = _PS("c", fontName="MSYH", fontSize=9, leading=13)

    buf = _io.BytesIO()
    doc = _Doc(buf, pagesize=_A4, leftMargin=2 * _cm, rightMargin=2 * _cm)
    els = [_P("桥梁病害检测报告", st_title),
           _P(f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}　·　桥梁病害检测平台", st_n),
           _P("一、基本信息", st_h)]
    info = [["文件名", filename], ["检测时间", created_at],
            ["检出病害", f"{len(detections)} 处"], ["模型", "YOLOv26-seg（桥梁病害 5 类）"]]
    t = _T([[ _P(a, st_c), _P(b, st_c)] for a, b in info], colWidths=[3.2 * _cm, 11.8 * _cm])
    t.setStyle(_TS([("GRID", (0, 0), (-1, -1), 0.5, _colors.grey),
                    ("BACKGROUND", (0, 0), (0, -1), _colors.HexColor("#f0f0f0"))]))
    els += [t, _P("二、标注结果图", st_h)]
    im = _Img(_io.BytesIO(ann_jpg))
    im._restrictSize(15 * _cm, 11 * _cm)
    els += [im, _P("三、病害明细", st_h)]
    if detections:
        head = ["序号", "类别", "置信度", "量化测量", "病害等级"]
        data = [[_P(x, st_c) for x in head]]
        for r in _report_rows(detections):
            data.append([_P(x, st_c) for x in r])
        t2 = _T(data, colWidths=[1.2 * _cm, 2.2 * _cm, 2.0 * _cm, 5.6 * _cm, 4.0 * _cm], repeatRows=1)
        t2.setStyle(_TS([("GRID", (0, 0), (-1, -1), 0.5, _colors.grey),
                         ("BACKGROUND", (0, 0), (-1, 0), _colors.HexColor("#2dd4bf")),
                         ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
        els.append(t2)
    else:
        els.append(_P("本次检测未发现病害。", st_n))
    els.append(_Sp(1, 0.4 * _cm))
    els.append(_P("注：量化测量为像素级估算（未做米制标定）；等级依据内置评估规则，仅供参考。", st_n))
    doc.build(els)
    return buf.getvalue()


@app.get("/report/{det_id}")
def export_report(det_id: int, format: str = "docx", user: dict = Depends(auth.get_current_user)):
    if format not in ("docx", "pdf"):
        raise HTTPException(status_code=400, detail="format 仅支持 docx/pdf")
    conn = db.get_db()
    row = conn.execute("SELECT * FROM detections WHERE id = ? AND user_id = ?",
                       (det_id, user["id"])).fetchone()
    conn.close()
    if row is None:
        raise HTTPException(status_code=404, detail="记录不存在")
    img_p = UPLOAD_DIR / row["image_path"]
    if not img_p.exists():
        raise HTTPException(status_code=404, detail="检测图片已删除")
    img = cv2.imdecode(np.fromfile(str(img_p), dtype=np.uint8), cv2.IMREAD_COLOR)
    if img is None:
        raise HTTPException(status_code=500, detail="检测图片损坏")
    detections = __import__("json").loads(row["result_json"])
    ann = _annotate_report(img, detections)
    ok, buf = cv2.imencode(".jpg", ann, [cv2.IMWRITE_JPEG_QUALITY, 90])
    ann_jpg = buf.tobytes() if ok else img_p.read_bytes()
    if format == "pdf":
        data = _export_pdf(row["filename"], row["created_at"], detections, ann_jpg)
        media, ext = "application/pdf", "pdf"
    else:
        data = _export_docx(row["filename"], row["created_at"], detections, ann_jpg)
        media, ext = ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx")
    return Response(content=data, media_type=media,
                    headers={"Content-Disposition": f'attachment; filename="report_{det_id}_{datetime.now():%Y%m%d}.{ext}"'})


# ============================================================
# 统计分析 + 数据集信息
# ============================================================
@app.get("/stats")
def stats(user: dict = Depends(auth.get_current_user)):
    conn = db.get_db()
    rows = conn.execute(
        "SELECT result_json, count, created_at FROM detections WHERE user_id = ?",
        (user["id"],)).fetchall()
    conn.close()

    total_images = len(rows)
    total_boxes = sum(r["count"] for r in rows)
    class_count = {c["id"]: 0 for c in DATASET_CLASSES}
    daily = {}
    for r in rows:
        dets = __import__("json").loads(r["result_json"])
        for d in dets:
            if d.get("rule_hint"):
                continue  # 规则提示不计入成绩统计
            class_count[d["class_id"]] = class_count.get(d["class_id"], 0) + 1
        day = r["created_at"][:10]
        daily[day] = daily.get(day, 0) + 1

    return {
        "total_images": total_images,
        "total_boxes": total_boxes,
        "class_distribution": [{"class": c["name"], "zh": c["zh"], "color": c["color"],
                                "count": class_count[c["id"]]} for c in DATASET_CLASSES],
        "daily": [{"date": k, "count": v} for k, v in sorted(daily.items())],
    }


@app.get("/dataset")
def dataset_info():
    return {"classes": DATASET_CLASSES, "splits": DATASET_SPLITS,
            "total_images": sum(DATASET_SPLITS.values()),
            "total_boxes": sum(c["boxes"] for c in DATASET_CLASSES)}


BASELINE_DIR = Path(__file__).resolve().parent / "data" / "baselines"
BASELINE_META = [
    ("honest_s", "第一轮 s（部署版）", "闭卷测试集 1665 张"),
    ("honest_distill", "第三轮 distill", "闭卷测试集 1665 张"),
    ("honest_m", "第二轮 m（教师）", "闭卷测试集 1665 张"),
    ("clean_s", "第一轮 s", "清洗后全测试集 2406 张（含旧模型开卷红利，仅供参考）"),
    ("clean_distill", "第三轮 distill", "清洗后全测试集 2406 张（含旧模型开卷红利，仅供参考）"),
    ("clean_m", "第二轮 m", "清洗后全测试集 2406 张（含旧模型开卷红利，仅供参考）"),
]


@app.get("/baselines")
def baselines(user: dict = Depends(auth.get_current_user)):
    """预存模型评测基线（mask mAP50 总体 + 分项，供对比页与论文截图）"""
    import json as _json

    out = []
    for tag, label, note in BASELINE_META:
        p = BASELINE_DIR / f"{tag}.json"
        if not p.exists():
            continue
        try:
            d = _json.loads(p.read_text(encoding="utf-8"))
            out.append({"tag": tag, "label": label, "note": note,
                        "mask_all": d["mask"]["mAP50"], "box_all": d["box"]["mAP50"],
                        "mask_per_class": d["mask"]["per_class_mAP50"],
                        "box_per_class": d["box"]["per_class_mAP50"]})
        except Exception:
            continue
    return {"baselines": out, "metric": "mAP50", "classes": [c["name"] for c in DATASET_CLASSES]}


@app.get("/health")
def health():
    return {"status": "ok", "model_loaded": _model is not None, "model_path": str(MODEL_PATH)}


# 已保存检测图的静态访问（uuid 文件名，不可猜测）
app.mount("/image", StaticFiles(directory=str(UPLOAD_DIR)), name="image")

# 融合演示场景图（随仓库的仿真样本）
app.mount("/fusion_img", StaticFiles(directory=str(FUSION_DIR)), name="fusion_img")

# 演示素材库图片/视频的静态访问
app.mount("/lib_img", StaticFiles(directory=str(LIB_DIR)), name="lib_img")
app.mount("/lib_video", StaticFiles(directory=str(LIB_VIDEO_DIR)), name="lib_video")

# 标注视频的静态访问
app.mount("/video", StaticFiles(directory=str(VIDEO_DIR)), name="video")

# 前端静态页（挂在最后，让 API 路由优先匹配）
FRONTEND_DIR = Path(__file__).resolve().parent.parent / "frontend"
app.mount("/", StaticFiles(directory=str(FRONTEND_DIR), html=True), name="frontend")
